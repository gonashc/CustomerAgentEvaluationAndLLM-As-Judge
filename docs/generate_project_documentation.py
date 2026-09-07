"""Generate the Word project documentation from the latest evaluation artifacts."""

from __future__ import annotations

import csv
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Customer_Support_Agent_Evaluation_Project_Documentation.docx"
CATEGORIES = [
    "order_status",
    "refund_request",
    "product_issue",
    "account_help",
    "other",
]

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "00A6A6"
LIGHT_BLUE = "DCE6F1"
LIGHT_TEAL = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
DARK_GRAY = "404040"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"


def load_rows(filename: str) -> list[dict[str, str]]:
    with (ROOT / filename).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def summarize(rows: list[dict[str, str]]) -> dict[str, object]:
    correct = sum(r["true_category"] == r["predicted_category"] for r in rows)
    failures = [r for r in rows if r["true_category"] != r["predicted_category"]]
    return {
        "rows": len(rows),
        "correct": correct,
        "failures": len(failures),
        "accuracy": correct / len(rows) if rows else 0,
        "patterns": Counter(
            (r["true_category"], r["predicted_category"]) for r in failures
        ),
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9E2F3", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, label in enumerate(headers):
        cell = header.cells[idx]
        cell.text = str(label)
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        if widths:
            cell.width = widths[idx]
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
            if row_idx % 2:
                set_cell_shading(cells[idx], LIGHT_GRAY)
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if widths:
                cells[idx].width = widths[idx]
    return table


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.right_indent = Inches(0.3)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F7F9FB")
    p_pr.append(shading)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(32, 55, 80)


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=BLUE, size="10")
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title + "\n")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(23, 54, 93)
    body_run = paragraph.add_run(text)
    body_run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_status_paragraph(doc: Document, label: str, text: str, color: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(label + "  ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    paragraph.add_run(text)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in [
        ("Title", 30, NAVY),
        ("Subtitle", 15, BLUE),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11, TEAL),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "CUSTOMER SUPPORT AGENT EVALUATION  |  PROJECT DOCUMENTATION"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        add_page_number(section.footer.paragraphs[0])


def build_document() -> Path:
    v1 = load_rows("results_v1.csv")
    v2 = load_rows("results_v2.csv")
    s1 = summarize(v1)
    s2 = summarize(v2)

    v2_by_id = {row["id"]: row for row in v2}
    comparisons = []
    for row in v1:
        newer = v2_by_id[row["id"]]
        c1 = row["true_category"] == row["predicted_category"]
        c2 = newer["true_category"] == newer["predicted_category"]
        comparisons.append((c1, c2, row["predicted_category"] != newer["predicted_category"]))
    wins = sum(not c1 and c2 for c1, c2, _ in comparisons)
    regressions = sum(c1 and not c2 for c1, c2, _ in comparisons)
    changed = sum(changed for _, _, changed in comparisons)
    improvement = float(s2["accuracy"]) - float(s1["accuracy"])

    doc = Document()
    style_document(doc)
    core = doc.core_properties
    core.title = "Customer Support Agent Evaluation — Project Documentation"
    core.subject = "LLM classifier evaluation, prompt iteration, and LangSmith observability"
    core.author = "Project Team"
    core.keywords = "LLM evaluation, LangGraph, LangSmith, OpenTelemetry, customer support"
    core.comments = "Generated from the repository's latest saved evaluation artifacts."

    # Cover page
    doc.add_paragraph("CUSTOMER SUPPORT AGENT EVALUATION", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Project Documentation and Outcome Report", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run("\n")
    cover = doc.add_table(rows=4, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.style = "Table Grid"
    cover_data = [
        ("Project", "CustomerAgentEvaluationAndLLM-As-Judge"),
        ("Report version", "1.0"),
        ("Prepared", date(2026, 9, 7).strftime("%d %B %Y")),
        ("Overall status", "Core evaluation and tracing completed; optional LLM judge not run"),
    ]
    for row, (label, value) in zip(cover.rows, cover_data):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], NAVY)
        set_cell_border(row.cells[0])
        set_cell_border(row.cells[1])
        for run in row.cells[0].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().add_run("\n")
    add_callout(
        doc,
        "Outcome at a glance",
        f"The focused prompt iteration raised accuracy from {float(s1['accuracy']):.0%} "
        f"to {float(s2['accuracy']):.0%} on the same 100-ticket dataset. The comparison "
        f"contains {wins} wrong-to-right corrections, {regressions} regressions, and "
        f"{int(s2['failures'])} remaining error. LangSmith ingestion and evaluation metadata "
        "were independently verified.",
        GREEN,
    )
    doc.add_paragraph().add_run("\n")
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Confidentiality note: API keys are intentionally excluded from this report.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    doc.add_page_break()

    # Contents
    doc.add_heading("Document map", level=1)
    contents = [
        "1. Executive summary",
        "2. Objectives and scope",
        "3. Solution architecture",
        "4. Environment and configuration",
        "5. Evaluation workflow",
        "6. Results and failure analysis",
        "7. Focused prompt revision (Part 3)",
        "8. LangSmith and OpenTelemetry observability (Part 6)",
        "9. Troubleshooting record",
        "10. Optional work and current status",
        "11. Reproducibility, security, and limitations",
        "12. Handoff checklist and conclusion",
    ]
    for item in contents:
        add_bullet(doc, item)
    doc.add_page_break()

    # 1
    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        "This project evaluates a LangGraph-based e-commerce customer-support routing agent. "
        "Given a ticket, the agent assigns exactly one of five routes: order status, refund "
        "request, product issue, account help, or other. The work demonstrates an end-to-end "
        "evaluation loop: construct a labeled dataset, run a baseline prompt, inspect failure "
        "patterns, make one focused prompt change, rerun on the same records, compare regressions, "
        "and inspect individual executions in LangSmith through OpenTelemetry."
    )
    add_table(
        doc,
        ["Measure", "Baseline", "Improved", "Change"],
        [
            ["Accuracy", f"{float(s1['accuracy']):.0%}", f"{float(s2['accuracy']):.0%}", f"{improvement:+.0%}"],
            ["Correct predictions", f"{s1['correct']} / {s1['rows']}", f"{s2['correct']} / {s2['rows']}", f"+{int(s2['correct']) - int(s1['correct'])}"],
            ["Classifier failures", s1["failures"], s2["failures"], f"-{int(s1['failures']) - int(s2['failures'])}"],
            ["Regression analysis", "—", f"{wins} wins / {regressions} regressions", f"{changed} changed routes"],
        ],
    )
    doc.add_paragraph(
        "The result meets the core learning objective: a measurable prompt change improved the "
        "routing behavior without harming previously correct examples in the latest exported run. "
        "The optional LLM-as-a-judge activity was intentionally left unexecuted."
    )

    # 2
    doc.add_heading("2. Objectives and scope", level=1)
    doc.add_heading("Primary objectives", level=2)
    for text in [
        "Build a structured-output classifier with LangGraph and an OpenAI chat model.",
        "Create a balanced, labeled evaluation dataset containing realistic and ambiguous tickets.",
        "Measure baseline accuracy, per-class performance, and confusion patterns.",
        "Select one business-relevant failure category and apply one focused prompt modification.",
        "Run a regression comparison on the same ticket records.",
        "Record ticket-level traces and evaluation metadata in LangSmith using OpenTelemetry.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Scope boundaries", level=2)
    doc.add_paragraph(
        "The work is an evaluation prototype rather than a production help-desk integration. It "
        "does not update customer orders, issue refunds, authenticate users, or connect to a live "
        "ticketing platform. The LLM-as-a-judge exercise and its associated alignment step are optional "
        "and were not required for completion of the classifier and trace-inspection work."
    )

    # 3
    doc.add_heading("3. Solution architecture", level=1)
    add_code(
        doc,
        "Labeled ticket → Prompt template → ChatOpenAI structured output\n"
        "              → LangGraph classifier → Prediction + reasoning\n"
        "              → pandas evaluation → CSV / spreadsheet review\n"
        "              → OpenTelemetry span → LangSmith project",
    )
    add_table(
        doc,
        ["Component", "Responsibility"],
        [
            ["LangGraph", "Compiles the classification workflow and exposes a consistent invoke interface."],
            ["ChatOpenAI", "Produces a typed category and a short routing reason."],
            ["Pydantic", "Constrains output to the five permitted category values."],
            ["pandas / scikit-learn", "Builds result tables and computes accuracy, reports, and confusion matrices."],
            ["OpenTelemetry", "Creates one ticket-level span and attaches evaluation metadata."],
            ["LangSmith", "Stores and displays traces for failure and regression inspection."],
        ],
    )

    # 4
    doc.add_heading("4. Environment and configuration", level=1)
    doc.add_heading("Runtime", level=2)
    add_table(
        doc,
        ["Item", "Configured value"],
        [
            ["Operating environment", "Windows / VS Code Jupyter"],
            ["Virtual environment", ".venv-x64"],
            ["Python", "3.12.13, 64-bit AMD64"],
            ["Model used", "gpt-4o-mini"],
            ["Classifier temperature", "0"],
            ["Synthetic expansion temperature", "0.9"],
        ],
    )
    doc.add_heading("Principal package versions", level=2)
    add_table(
        doc,
        ["Package", "Version"],
        [
            ["langgraph", "1.2.11"],
            ["langsmith", "0.12.2"],
            ["langchain-openai", "1.6.0"],
            ["pandas", "3.0.5"],
            ["scikit-learn", "1.9.0"],
            ["pydantic", "2.13.5"],
            ["opentelemetry-sdk", "1.44.0"],
        ],
    )
    doc.add_heading("Secret and tracing configuration", level=2)
    doc.add_paragraph(
        "Configuration is stored in the project-local .env file, which is excluded from source "
        "control. The notebook loads it with load_dotenv(override=True) so replacement keys supersede "
        "values cached by an earlier kernel session. Only variable names are documented here:"
    )
    add_code(
        doc,
        "OPENAI_API_KEY\nLANGSMITH_API_KEY\nLANGSMITH_TRACING=true\n"
        "LANGSMITH_PROJECT=CustomerAgentEvaluationAndLLM-As-Judge\n"
        "LANGSMITH_ENDPOINT=https://api.smith.langchain.com\n"
        "LANGSMITH_OTEL_ENABLED=true\n"
        "OTEL_EXPORTER_OTLP_ENDPOINT=https://api.smith.langchain.com/otel/v1/traces\n"
        "OTEL_EXPORTER_OTLP_TRACES_ENDPOINT=https://api.smith.langchain.com/otel/v1/traces\n"
        "OTEL_EXPORTER_OTLP_PROTOCOL=http/protobuf",
    )

    # 5
    doc.add_heading("5. Evaluation workflow", level=1)
    workflow = [
        "Install dependencies in the 64-bit project virtual environment.",
        "Load API and tracing configuration before importing and invoking model-backed components.",
        "Define the five routing categories and create 25 trusted seed tickets.",
        "Generate three paraphrases per seed, producing 100 tickets total and 20 per category.",
        "Build and smoke-test the baseline LangGraph classifier.",
        "Create a ticket-level OpenTelemetry tracer and evaluation metadata schema.",
        "Run the baseline classifier over all 100 tickets and export results_v1.csv.",
        "Inspect failures and group recurring confusion patterns.",
        "Document the chosen business-priority failure and add one disambiguation rule.",
        "Run the improved classifier on the same tickets and export results_v2.csv.",
        "Compare changed predictions, wrong-to-right wins, and right-to-wrong regressions.",
        "Inspect the corresponding baseline and improved traces in LangSmith.",
    ]
    for step in workflow:
        add_number(doc, step)
    doc.add_heading("Dataset composition", level=2)
    add_table(
        doc,
        ["Category", "Tickets", "Typical intent"],
        [
            ["order_status", 20, "Tracking, delivery ETA, missing or undelivered orders"],
            ["refund_request", 20, "Returns, cancellations, refund status, price adjustment"],
            ["product_issue", 20, "Damaged, defective, wrong, incomplete, or misleading products"],
            ["account_help", 20, "Login, password, profile, payment method, and account security"],
            ["other", 20, "General questions, availability, discounts, feedback, and sales"],
        ],
    )

    # 6
    doc.add_heading("6. Results and failure analysis", level=1)
    doc.add_heading("Baseline confusion patterns", level=2)
    pattern_rows = [
        [true, predicted, count]
        for (true, predicted), count in sorted(
            s1["patterns"].items(), key=lambda item: (-item[1], item[0])
        )
    ]
    add_table(doc, ["True category", "Predicted category", "Count"], pattern_rows)
    doc.add_paragraph(
        "The dominant error mechanism was remedy language overriding the underlying operational "
        "problem. Tickets about missing orders or damaged products often mentioned money back, and "
        "the baseline classifier routed them to refund_request. A subtler not-as-described product "
        "case was routed to other."
    )
    doc.add_heading("Per-category recall", level=2)
    add_table(
        doc,
        ["Category", "Baseline", "Improved"],
        [
            ["order_status", "80% (16/20)", "100% (20/20)"],
            ["refund_request", "100% (20/20)", "100% (20/20)"],
            ["product_issue", "75% (15/20)", "95% (19/20)"],
            ["account_help", "100% (20/20)", "100% (20/20)"],
            ["other", "100% (20/20)", "100% (20/20)"],
        ],
    )
    doc.add_heading("Improved-run outcome", level=2)
    add_status_paragraph(doc, "PASS", f"Accuracy improved by {improvement:+.0%} to {float(s2['accuracy']):.0%}.", "548235")
    add_status_paragraph(doc, "PASS", f"The comparison recorded {wins} wins and {regressions} regressions.", "548235")
    add_status_paragraph(doc, "OPEN", "One product_issue ticket remains classified as other.", "BF7000")

    # 7
    doc.add_heading("7. Focused prompt revision (Part 3)", level=1)
    doc.add_heading("Selected failure category", level=2)
    doc.add_paragraph("Product issues misrouted as refund requests.")
    doc.add_heading("Business rationale", level=2)
    doc.add_paragraph(
        "Damaged, defective, wrong, incomplete, or misleading products require product-support and "
        "quality-investigation workflows. Routing them only to refund processing can conceal supplier, "
        "inventory, packaging, and safety trends."
    )
    doc.add_heading("Focused disambiguation rule", level=2)
    add_code(
        doc,
        "If an item is damaged, wrong, defective, missing an advertised capability,\n"
        "or not as described, classify it as product_issue even when the customer\n"
        "also requests a refund.",
    )
    doc.add_paragraph(
        "The implementation derives IMPROVED_PROMPT from CLASSIFIER_PROMPT and inserts this single "
        "rule before the return instruction. It does not replace the rest of the baseline prompt. "
        "The complete 100-row rerun serves as a regression check and is stronger than retesting only "
        "one or two examples."
    )
    add_callout(
        doc,
        "Interpretation caution",
        "The latest comparison contains improvements outside the specifically targeted product-issue "
        "pattern. Because synthetic expansion uses a nonzero temperature and hosted model execution is "
        "not perfectly deterministic, the full eight-point gain should not be attributed solely to the "
        "new sentence without repeated or frozen-dataset trials.",
        AMBER,
    )

    # 8
    doc.add_heading("8. LangSmith and OpenTelemetry observability (Part 6)", level=1)
    doc.add_paragraph(
        "The notebook emits a parent span named customer_support.classify_ticket for each evaluation "
        "record. The span is sent to the US LangSmith endpoint and supplements the nested LangGraph and "
        "model traces with evaluation-specific fields."
    )
    add_table(
        doc,
        ["Verified item", "Evidence"],
        [
            ["LangSmith project", "CustomerAgentEvaluationAndLLM-As-Judge exists and is queryable."],
            ["Ticket-level trace ingestion", "400 customer_support.classify_ticket roots across repeated baseline/improved runs."],
            ["Trace health", "0 root trace errors at verification time."],
            ["Record linkage", "example_id is present in latest trace metadata."],
            ["Run separation", "run_name and prompt_version are present."],
            ["Evaluation inspection", "true_category, predicted_category, correct, and reasoning are present."],
        ],
    )
    doc.add_heading("Metadata attached to each ticket span", level=2)
    add_code(
        doc,
        "run_name | prompt_version | example_id | true_category |\n"
        "predicted_category | correct | reasoning",
    )
    doc.add_paragraph(
        "The LangSmith metadata helper prefixes the attributes in the format expected by LangSmith. "
        "This allows a reviewer to find a CSV row by example ID, compare v1 and v2 behavior, and inspect "
        "the model's rationale without relying on aggregate accuracy alone. Part 6 is technically complete."
    )

    # 9
    doc.add_heading("9. Troubleshooting record", level=1)
    add_table(
        doc,
        ["Issue", "Root cause", "Resolution", "Outcome"],
        [
            ["Environment/interpreter mismatch", "The original environment was not the intended 64-bit runtime.", "Created and selected .venv-x64 using Python 3.12.13 AMD64.", "Required packages import successfully."],
            ["pip unavailable in new environment", "The virtual environment initially lacked the pip module.", "Bootstrapped pip with python -m ensurepip --upgrade.", "Notebook dependency installation became available."],
            ["OpenAI 401 invalid_api_key", "The key stored in .env was rejected by the OpenAI API.", "Replaced the key, loaded .env with override=True, and restarted the kernel.", "Dataset generation and classification calls completed."],
            ["LangSmith spans returned 404", "The general OTLP URL ended at /otel rather than the full traces path.", "Set both OTLP endpoint variables to /otel/v1/traces and restarted the kernel.", "Diagnostic and evaluation traces reached LangSmith."],
            ["Project naming mismatch", "The notebook and .env used different project names.", "Standardized on CustomerAgentEvaluationAndLLM-As-Judge.", "Traces are collected in the intended project."],
            ["Evaluation metadata absent", "Raw eval.* attributes were not exposed as LangSmith metadata.", "Used set_langsmith_metadata_attribute for seven evaluation fields.", "Latest ticket traces contain all required metadata."],
            ["IndentationError in Section 6", "The inserted metadata dictionary and loop were dedented outside the active span block.", "Aligned the dictionary, loop, and span attributes inside the with block.", "The saved source passes Python syntax validation."],
            ["Notebook output versus CSV mismatch", "Notebook outputs and generated CSVs came from different saved/rerun states.", "Use timestamped CSVs as the latest measurement and save/clear stale notebook output before submission.", "Latest exported results are internally consistent."],
        ],
    )

    # 10
    doc.add_heading("10. Optional work and current status", level=1)
    add_table(
        doc,
        ["Work item", "Status", "Notes"],
        [
            ["Core classifier evaluation", "Complete", "Baseline and improved results exported for the same 100 tickets."],
            ["Part 3 — focused prompt tweak", "Complete", "One documented business failure and one disambiguation rule; full regression comparison completed."],
            ["Part 4 — LLM-as-a-judge", "Not run (optional)", "Would require completed human failure labels and a fixed target category before judging each designated failure row."],
            ["Part 5 — optional follow-on", "Skipped", "Skipping optional work does not block trace inspection."],
            ["Part 6 — LangSmith trace inspection", "Complete", "Trace delivery and required metadata were verified programmatically."],
        ],
    )
    doc.add_heading("Requirements if Part 4 is later attempted", level=2)
    for text in [
        "Complete the Step 3 human failure label for every designated classifier-failure row.",
        "Choose exactly one target failure category before invoking the judge.",
        "Define ground truth as an exact match between each human label and that target category.",
        "Ask the judge for a binary TRUE/FALSE prediction plus a short reason without exposing the human label.",
        "Calculate agreement between judge prediction and human ground truth.",
    ]:
        add_bullet(doc, text)
    add_callout(
        doc,
        "Dataset alignment warning",
        "The optional instructions reference 29 fixed classifier-failure rows. The latest generated "
        f"baseline CSV contains {s1['failures']} failures and blank validator/failure-category fields. "
        "If Part 4 is performed, use the designated fixed spreadsheet rows rather than silently "
        "substituting the stochastic notebook failures.",
        AMBER,
    )

    # 11
    doc.add_heading("11. Reproducibility, security, and limitations", level=1)
    doc.add_heading("Reproducibility", level=2)
    for text in [
        "The same ticket texts were used for the latest v1 and v2 CSV comparison.",
        "Synthetic paraphrase generation uses temperature 0.9, so a full notebook rerun creates a different dataset.",
        "Freeze and version the generated ticket dataset before repeated prompt experiments.",
        "Record model name, package versions, prompt version, project name, and execution timestamp for each run.",
        "Use repeated trials or a larger fixed evaluation set before interpreting small accuracy differences causally.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Security", level=2)
    for text in [
        "Keep .env, API keys, and copied secret values outside source control and submission artifacts.",
        "Revoke and replace a key immediately if it is exposed.",
        "Avoid printing keys in notebook output; report only presence or masked diagnostics.",
        "Use the correct LangSmith regional endpoint for the workspace associated with the key.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Evaluation limitations", level=2)
    doc.add_paragraph(
        "The dataset is synthetic, small, balanced, and deliberately designed around five labels. "
        "Production traffic is likely imbalanced and may include multilingual content, multiple intents, "
        "policy-sensitive cases, and category drift. Accuracy should therefore be treated as evidence for "
        "the exercise—not a production service-level guarantee."
    )

    # 12
    doc.add_heading("12. Handoff checklist and conclusion", level=1)
    checklist = [
        ("Complete", "100-ticket baseline and improved CSV files are present."),
        ("Complete", "Focused Part 3 choice and prompt rule are documented."),
        ("Complete", "Latest comparison shows 91% → 99%, eight wins, and zero regressions."),
        ("Complete", "LangSmith has ticket-level traces with the seven required metadata fields."),
        ("Complete", "Historical error outputs were cleared from the release notebook."),
        ("Complete", "README metrics match the latest exported 91%/99% run."),
        ("Optional", "Populate validator comments/failure categories if human-review artifacts are required."),
        ("Optional", "Run the LLM-as-a-judge activity only against the designated labeled failure rows."),
    ]
    add_table(doc, ["Status", "Checklist item"], [[a, b] for a, b in checklist])
    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph(
        "The project successfully demonstrates a complete classifier-evaluation and observability loop. "
        "The agent processed a balanced 100-ticket dataset, the baseline weaknesses were measured, one "
        "business-relevant prompt rule was applied, and the improved run reached 99% accuracy without a "
        "regression in the latest export. LangSmith now stores the ticket-level executions with sufficient "
        "metadata to connect traces back to evaluation rows. Remaining work is limited to submission "
        "cleanup, reproducibility hardening, and explicitly optional evaluation activities."
    )

    # Appendix
    doc.add_page_break()
    doc.add_heading("Appendix A — Repository deliverables", level=1)
    add_table(
        doc,
        ["Artifact", "Purpose"],
        [
            ["week4_customer_support_evals.ipynb", "Executable classifier, evaluation, prompt iteration, and trace instrumentation."],
            ["results_v1.csv", "Latest baseline predictions, reasoning, correctness, and review columns."],
            ["results_v2.csv", "Latest focused-prompt predictions and evaluation records."],
            ["Week 4_ AI Evals (E-Commerce Customer Support Agent).xlsx", "Human review and evaluation workbook."],
            ["docs/opentelemetry_integration_one_pager.md", "Technical overview and troubleshooting for trace integration."],
            ["README.md", "Repository summary, setup instructions, and headline results."],
            ["docs/Customer_Support_Agent_Evaluation_Project_Documentation.docx", "This consolidated project report."],
        ],
    )
    doc.add_heading("Appendix B — Recommended clean execution order", level=1)
    for text in [
        "Select .venv-x64 as the Jupyter kernel.",
        "Confirm .env is saved and contains current keys, project, and full OTLP traces endpoint.",
        "Restart the kernel so no stale exporter or cached key remains.",
        "Run configuration before any LangChain or OpenTelemetry-backed model invocation.",
        "Run category definition, ticket generation, classifier build, and tracer definition.",
        "Run baseline classification and evaluation, then export v1.",
        "Run the focused improved prompt on the same in-memory tickets, evaluate, and export v2.",
        "Run the comparison cell and save the notebook.",
        "Refresh LangSmith and inspect traces by example_id, run_name, and prompt_version.",
    ]:
        add_number(doc, text)

    # Make a clean continuation section so Word has a consistent final footer.
    final_section = doc.add_section(WD_SECTION.CONTINUOUS)
    final_section.header.is_linked_to_previous = True
    final_section.footer.is_linked_to_previous = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
